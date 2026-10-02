import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from doc_helpers import (
    apply_page_border, apply_header, apply_footer,
    set_cell_border, set_cell_shading,
    add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_bullet_p, add_image_figure
)

ASSETS_DIR = r"C:\xampp\htdocs\stock_market\extracted_assets"
OUTPUT_DOCX = r"C:\xampp\htdocs\stock_market\Stock_Market_Application_Project_Report.docx"
DOWNLOADS_DOCX = r"C:\Users\Sneh Patel\Downloads\Stock_Market_Application_Project_Report.docx"

doc = docx.Document()

# ----------------- PAGE SETUP -----------------
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.4)
section.footer_distance = Inches(0.4)

apply_page_border(section)
apply_header(section, "Stock Market Web Application")

team_members = [
    "SNEH PATEL (2303031080112)",
    "BHUMI PATEL (2303031080093)",
    "SATYAM PATEL (2303031080073)",
    "BHAVIN RANA (2303031080243)"
]
apply_footer(section, team_members)

# ==========================================
# 1. TITLE / COVER PAGE
# ==========================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("In partial fulfillment for the award of the degree of")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("BACHELOR OF TECHNOLOGY")
r.font.name = "Times New Roman"
r.font.size = Pt(15)
r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("in\nINFORMATION TECHNOLOGY")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.font.bold = True

logo_path = os.path.join(ASSETS_DIR, "stock_p1_img1_804x240.png")
if os.path.exists(logo_path):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(4)
    p_logo.paragraph_format.space_after = Pt(6)
    p_logo.add_run().add_picture(logo_path, width=Inches(3.2))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("DEPARTMENT OF INFORMATION TECHNOLOGY,\nPARUL INSTITUTE OF ENGINEERING AND TECHNOLOGY,\nPARUL UNIVERSITY,\nVADODARA, GUJARAT")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Stock Market Application")
r.font.name = "Times New Roman"
r.font.size = Pt(22)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(12)
r = p.add_run("Web-Service to BUY & SELL Stocks")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Submitted by:")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.bold = True

students = [
    "MR. SNEH JITENDRAKUMAR PATEL (2303031080112)",
    "MS. BHUMI NARENDRABHAI PATEL (2303031080093)",
    "MR. SATYAM DATTUBHAI PATEL (2303031080073)",
    "MR. BHAVIN ANURAGBHAI RANA (2303031080243)"
]
for s in students:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(s)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Under the Guidance of")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Assistant Professor MS. SONALI KORI")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("[2025-2026]")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

# ==========================================
# 2. CERTIFICATE PAGE
# ==========================================
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(16)
r = p.add_run("CERTIFICATE")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

cert_text = (
    'This is to certify that the Project Report entitled, "Stock Market Application: Web-service to BUY & SELL Stocks" '
    'submitted by "Sneh Jitendrakumar Patel, Bhumi Narendrabhai Patel, Satyam Dattubhai Patel, Bhavin Anuragbhai Rana" '
    'to Parul University, Vadodara, Gujarat, is a record of bonafide Project work carried out by them under my supervision '
    'and guidance, and is worthy of consideration for the award of the degree of Bachelor of Technology in '
    'Information Technology of the University.'
)
add_body_p(doc, cert_text, space_after=18)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Date:  ____________________\nPlace: Vadodara, Gujarat")
r.font.name = "Times New Roman"
r.font.size = Pt(11)

sig_table = doc.add_table(rows=2, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in sig_table.rows:
    row.cells[0].width = Inches(3.5)
    row.cells[1].width = Inches(3.5)

cell_00 = sig_table.rows[0].cells[0]
p = cell_00.paragraphs[0]
p.paragraph_format.space_before = Pt(32)
r = p.add_run("______________________________\nSupervisor\nMs. Sonali Kori\nAssistant Professor")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

cell_01 = sig_table.rows[0].cells[1]
p = cell_01.paragraphs[0]
p.paragraph_format.space_before = Pt(32)
r = p.add_run("______________________________\nProject Coordinator\nMs. Dhenuka Patel\nAssistant Professor")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

cell_10 = sig_table.rows[1].cells[0]
p = cell_10.paragraphs[0]
p.paragraph_format.space_before = Pt(36)
r = p.add_run("______________________________\nHead, Dept. of Information Technology\nDr. Pooja Sapra")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

cell_11 = sig_table.rows[1].cells[1]
p = cell_11.paragraphs[0]
p.paragraph_format.space_before = Pt(36)
r = p.add_run("______________________________\nExternal Supervisor\nName:\nDesignation:")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

# ==========================================
# 3. ACKNOWLEDGEMENT PAGE
# ==========================================
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(16)
r = p.add_run("Acknowledgement")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

add_body_p(doc, 
    "It is a matter of great pleasure for us to get this opportunity expressing our sincere sense of gratitude. "
    "Firstly, we would like to express our heartfelt thanks to Parul University for providing modern laboratory "
    "and computing facilities. As part of our academic study as students of the Department of Information Technology, "
    "we are required to undergo practical project development to obtain technical knowledge and broaden our "
    "practical engineering competencies.", space_after=12)

add_body_p(doc, 
    "First and foremost, we would like to express our special and humble thanks and gratitude to Dr. Pooja Sapra "
    "(H.O.D, Department of Information Technology) who has provided us such a cooperative, progressive, and "
    "enriching academic environment. Secondly, we express our deepest gratitude to our Project Guide, "
    "Assistant Professor Ms. Sonali Kori, who contributed her invaluable time, rigorous technical feedback, and continuous "
    "encouragement throughout the design, implementation, and testing phases of this Stock Market Web Application.", space_after=12)

add_body_p(doc,
    "We also extend our sincere thanks to all faculty members, laboratory staff, and peer reviewers who directly or "
    "indirectly contributed their technical insights to the completion of this project.", space_after=20)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Submitted By: -\n")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

for s in students:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(s)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True

# ==========================================
# 4. ABSTRACT PAGE
# ==========================================
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(16)
r = p.add_run("ABSTRACT")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

add_body_p(doc, 
    "The Stock Market Application Web-Service is a full-featured web-based simulation and trading platform "
    "engineered to replicate the core operational dynamics of modern stock exchanges. The primary objective "
    "of this project is to provide retail users, university students, and novice investors with an intuitive, risk-free "
    "environment where they can register, analyze real-time market data, execute simulated buy and sell orders, "
    "manage multi-asset portfolios, and track transaction histories with financial precision.", space_after=10)

add_body_p(doc,
    "The architecture is designed following a robust client-server model utilizing HTML5, modern CSS3, and JavaScript "
    "for dynamic user interface rendering, PHP 8 for backend business logic and transactional processing, and MySQL "
    "Relational Database Management System (RDBMS) for persistent, ACID-compliant data storage. Essential functional "
    "modules include session-based user authentication, Alpha Vantage and Marketstack financial API integration, dynamic "
    "stock search, portfolio balance updates, and Razorpay API integration for simulated fund deposits.", space_after=10)

add_body_p(doc,
    "This project demonstrates the practical implementation of relational database schema design, server-side scripting, "
    "asynchronous data retrieval via AJAX, transaction atomicity, and security hardening against common vulnerabilities "
    "such as SQL injection and Cross-Site Scripting (XSS). The Stock Market Application serves as an effective educational "
    "technology and establishes a scalable foundation for future algorithmic trading and machine learning market forecasting.", space_after=14)

p_kw = doc.add_paragraph()
p_kw.paragraph_format.space_before = Pt(8)
r_k = p_kw.add_run("Keywords: ")
r_k.font.name = "Times New Roman"
r_k.font.size = Pt(11)
r_k.font.bold = True
r_v = p_kw.add_run("Stock Market, Web Application, PHP, MySQL, Trading Simulation, Financial APIs, Razorpay API, Portfolio Management")
r_v.font.name = "Times New Roman"
r_v.font.size = Pt(11)
r_v.font.italic = True

# Save preliminary document
doc.save(OUTPUT_DOCX)
print("Stock Market report preliminaries generated.")
