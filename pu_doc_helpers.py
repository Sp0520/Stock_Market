import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ASSETS_DIR = r"C:\xampp\htdocs\stock_market\extracted_assets"
OUTPUT_DOCX = r"C:\xampp\htdocs\stock_market\Stock_Market_Application_Project_Report.docx"
DOWNLOADS_DOCX = r"C:\Users\Sneh Patel\Downloads\Stock_Market_Application_Project_Report.docx"

def create_element(name):
    return OxmlElement(name)

def set_cell_border(cell, top="D3D3D3", bottom="D3D3D3", left="D3D3D3", right="D3D3D3", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f"""
    <w:tcBorders {nsdecls('w')}>
        <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>
        <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>
        <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>
        <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>
    </w:tcBorders>
    """)
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def apply_page_border(section):
    sectPr = section._sectPr
    borders_xml = f"""
    <w:pgBorders {nsdecls('w')} w:offsetFrom="page">
        <w:top w:val="single" w:sz="8" w:space="18" w:color="1F4E79"/>
        <w:left w:val="single" w:sz="8" w:space="18" w:color="1F4E79"/>
        <w:bottom w:val="single" w:sz="8" w:space="18" w:color="1F4E79"/>
        <w:right w:val="single" w:sz="8" w:space="18" w:color="1F4E79"/>
    </w:pgBorders>
    """
    sectPr.append(parse_xml(borders_xml))

def apply_header(section, title_text="Stock Market Web Application | Parul University"):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    hrun = hp.add_run(title_text)
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(9)
    hrun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    pBdr_xml = f"""
    <w:pBdr {nsdecls('w')}>
        <w:bottom w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/>
    </w:pBdr>
    """
    hp._p.get_or_add_pPr().append(parse_xml(pBdr_xml))

def apply_footer(section, prepared_by_names):
    footer = section.footer
    fp = footer.paragraphs[0]
    foot_tbl = footer.add_table(rows=1, cols=2, width=Inches(6.0))
    foot_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in foot_tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f"""
        <w:tcBorders {nsdecls('w')}>
            <w:top w:val="single" w:sz="10" w:space="0" w:color="1F4E79"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
        """)
        tcPr.append(tcBorders)

    left_cell = foot_tbl.rows[0].cells[0]
    left_cell.width = Inches(4.5)
    lp = left_cell.paragraphs[0]
    lp.paragraph_format.space_before = Pt(3)
    lp.paragraph_format.space_after = Pt(0)
    lp.paragraph_format.line_spacing = 1.0
    
    p_run = lp.add_run("PREPARED BY: " + " | ".join(prepared_by_names))
    p_run.font.name = "Calibri"
    p_run.font.size = Pt(8)
    p_run.font.bold = True
    p_run.font.color.rgb = RGBColor(0x44, 0x55, 0x66)

    right_cell = foot_tbl.rows[0].cells[1]
    right_cell.width = Inches(1.5)
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = Pt(3)
    rp.paragraph_format.space_after = Pt(0)
    
    rrun = rp.add_run("Page ")
    rrun.font.name = "Calibri"
    rrun.font.size = Pt(8.5)
    rrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    rp._p.append(fldSimple)

def add_chapter_title(doc, text):
    """PU Guideline: Chapter Title: Arial Rounded MT Bold (Upper Case), size 16 (e.g. CHAPTER 1)"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    run.font.name = "Arial Rounded MT Bold"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return p

def add_heading_1(doc, text):
    """PU Guideline: Main Heading: Calibri Bold, size 12 (e.g. 1.1 Introduction)"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return p

def add_heading_2(doc, text):
    """PU Guideline: Sub Heading: Calibri Bold, size 11 (e.g 1.1.1 Dynamic Source Routing Protocol)"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_body_p(doc, text, bold_prefix="", space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """PU Guideline: Body text: Calibri, size 11, Justified, 1.5 line spacing"""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if bold_prefix:
        brun = p.add_run(bold_prefix)
        brun.font.name = "Calibri"
        brun.font.size = Pt(11)
        brun.font.bold = True
        brun.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_bullet_p(doc, text, bold_prefix="", space_after=4):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if bold_prefix:
        brun = p.add_run(bold_prefix)
        brun.font.name = "Calibri"
        brun.font.size = Pt(11)
        brun.font.bold = True
        brun.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_callout_box(doc, text, title=""):
    """Adds a stylish, shaded callout box for key viva takeaways and exam tips."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.rows[0].cells[0].width = Inches(5.8)
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, "F0F4F8")
    set_cell_border(cell, top="1F4E79", bottom="1F4E79", left="1F4E79", right="1F4E79", sz="8")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if title:
        rt = p.add_run(title + "\n")
        rt.font.name = "Calibri"
        rt.font.size = Pt(10.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_table_caption(doc, caption_text):
    """PU Guideline: Table caption: Calibri, size 10, Centre aligned, Decimal type notation (e.g. Table 2.2 etc.)"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(caption_text)
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return p

def add_image_figure(doc, img_path, caption, width=Inches(5.4)):
    """PU Guideline: Figure caption: Calibri, size 10, Centre aligned, Decimal type notation (e.g. Figure 1.2 etc.)"""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        crun = p_cap.add_run(caption)
        crun.font.name = "Calibri"
        crun.font.size = Pt(10)
        crun.font.bold = True
        crun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_equation_box(doc, eq_text, eq_number):
    """PU Guideline: Equations should also be numbered in decimal type notation within brackets."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.rows[0].cells[0].width = Inches(5.0)
    tbl.rows[0].cells[1].width = Inches(0.9)
    
    c0 = tbl.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(eq_text)
    r0.font.name = "Calibri"
    r0.font.size = Pt(11)
    r0.font.bold = True
    r0.font.italic = True
    
    c1 = tbl.rows[0].cells[1]
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p1.add_run(eq_number)
    r1.font.name = "Calibri"
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    
    for c in [c0, c1]:
        set_cell_border(c, top="FFFFFF", bottom="FFFFFF", left="FFFFFF", right="FFFFFF")
        set_cell_shading(c, "F9FBFD")

print("PU Helper module loaded successfully.")
