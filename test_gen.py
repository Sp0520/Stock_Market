import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.4)
section.footer_distance = Inches(0.4)

sectPr = section._sectPr
borders_xml = f"""
<w:pgBorders {nsdecls('w')} w:offsetFrom="page">
    <w:top w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
    <w:left w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
    <w:bottom w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
    <w:right w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
</w:pgBorders>
"""
sectPr.append(parse_xml(borders_xml))

# Test Header
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hrun = hp.add_run("Ride Buddy Web Application")
hrun.font.name = "Times New Roman"
hrun.font.size = Pt(9.5)
hrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

pBdr_xml = f"""
<w:pBdr {nsdecls('w')}>
    <w:bottom w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/>
</w:pBdr>
"""
hp._p.get_or_add_pPr().append(parse_xml(pBdr_xml))

# Test Footer
footer = section.footer
fp = footer.paragraphs[0]

# Add horizontal blue bar on footer
foot_tbl = footer.add_table(rows=1, cols=2, width=Inches(7.0))
foot_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell in foot_tbl.rows[0].cells:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f"""
    <w:tcBorders {nsdecls('w')}>
        <w:top w:val="single" w:sz="12" w:space="0" w:color="2B579A"/>
        <w:left w:val="none"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    """)
    tcPr.append(tcBorders)

left_cell = foot_tbl.rows[0].cells[0]
lp = left_cell.paragraphs[0]
lp.paragraph_format.space_before = Pt(4)
lp.paragraph_format.space_after = Pt(0)
lrun = lp.add_run("PREPARED BY: AKSHAR PATEL | OM DOBARIYA | DIVYRAJSINH THAKOR | DHRUVANSH MANIYA")
lrun.font.name = "Times New Roman"
lrun.font.size = Pt(7.5)
lrun.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

right_cell = foot_tbl.rows[0].cells[1]
rp = right_cell.paragraphs[0]
rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rp.paragraph_format.space_before = Pt(4)
rp.paragraph_format.space_after = Pt(0)
rrun = rp.add_run()
rrun.font.name = "Times New Roman"
rrun.font.size = Pt(8.5)
rrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# Add page number field to right cell
fldSimple = OxmlElement('w:fldSimple')
fldSimple.set(qn('w:instr'), 'PAGE')
rp._p.append(fldSimple)

doc.add_paragraph("Hello world, this is a test document.")
doc.save("test_out.docx")
print("Saved test_out.docx successfully")
