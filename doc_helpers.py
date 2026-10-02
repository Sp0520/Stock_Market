import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def apply_page_border(section):
    sectPr = section._sectPr
    borders_xml = f"""
    <w:pgBorders {nsdecls('w')} w:offsetFrom="page">
        <w:top w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
        <w:left w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
        <w:bottom w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
        <w:right w:val="single" w:sz="8" w:space="18" w:color="2B579A"/>
    </w:pgBorders>
    """
    sectPr.append(parse_xml(borders_xml))

def apply_header(section, title_text="Ride Buddy Web Application"):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    hrun = hp.add_run(title_text)
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(9.5)
    hrun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    pBdr_xml = f"""
    <w:pBdr {nsdecls('w')}>
        <w:bottom w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/>
    </w:pBdr>
    """
    hp._p.get_or_add_pPr().append(parse_xml(pBdr_xml))

def apply_footer(section, prepared_by_names):
    footer = section.footer
    fp = footer.paragraphs[0]
    
    # We create a 1-row, 2-col table spanning page width
    foot_tbl = footer.add_table(rows=1, cols=2, width=Inches(7.0))
    foot_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in foot_tbl.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f"""
        <w:tcBorders {nsdecls('w')}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="2B579A"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
        """)
        tcPr.append(tcBorders)

    left_cell = foot_tbl.rows[0].cells[0]
    left_cell.width = Inches(5.5)
    lp = left_cell.paragraphs[0]
    lp.paragraph_format.space_before = Pt(3)
    lp.paragraph_format.space_after = Pt(0)
    lp.paragraph_format.line_spacing = 1.05
    
    p_run = lp.add_run("PREPARED BY:\n" + "\n".join(prepared_by_names))
    p_run.font.name = "Times New Roman"
    p_run.font.size = Pt(7.5)
    p_run.font.bold = True
    p_run.font.color.rgb = RGBColor(0x55, 0x66, 0x77)

    right_cell = foot_tbl.rows[0].cells[1]
    right_cell.width = Inches(1.5)
    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_before = Pt(3)
    rp.paragraph_format.space_after = Pt(0)
    
    rrun = rp.add_run()
    rrun.font.name = "Times New Roman"
    rrun.font.size = Pt(8.5)
    rrun.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    rp._p.append(fldSimple)

def set_cell_border(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f"""
    <w:tcBorders {nsdecls('w')}>
        <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>
        <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>
        <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>
        <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>
    </w:tcBorders>
    """)
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Elegant Navy Blue
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return h

def add_body_p(doc, text, bold_prefix="", space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        brun = p.add_run(bold_prefix)
        brun.font.name = "Times New Roman"
        brun.font.size = Pt(11)
        brun.font.bold = True
        brun.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_bullet_p(doc, text, bold_prefix="", space_after=3):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        brun = p.add_run(bold_prefix)
        brun.font.name = "Times New Roman"
        brun.font.size = Pt(11)
        brun.font.bold = True
        brun.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_image_figure(doc, img_path, caption, width=Inches(5.5)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        crun = p_cap.add_run(caption)
        crun.font.name = "Times New Roman"
        crun.font.size = Pt(10)
        crun.font.bold = True
        crun.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

print("Helper functions defined successfully")
