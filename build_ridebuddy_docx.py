import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from doc_helpers import (
    apply_page_border, apply_header, apply_footer,
    set_cell_border, set_cell_shading,
    add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_bullet_p, add_image_figure
)

ASSETS_DIR = r"C:\xampp\htdocs\stock_market\extracted_assets"
OUTPUT_DOCX = r"C:\xampp\htdocs\stock_market\Ride_Buddy_Project_Report.docx"
DOWNLOADS_DOCX = r"C:\Users\Sneh Patel\Downloads\Ride_Buddy_Project_Report.docx"

doc = docx.Document()

# Page Setup
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.4)
section.footer_distance = Inches(0.4)

apply_page_border(section)
apply_header(section, "Ride Buddy Web Application")

team_members = [
    "AKSHAR PATEL (2303031080220)",
    "OM DOBARIYA (2303031080039)",
    "DIVYRAJSINH THAKOR (2303031080249)",
    "DHRUVANSH MANIYA (2303031080212)"
]
apply_footer(section, team_members)

# ==========================================
# 1. TITLE / COVER PAGE
# ==========================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("In partial fulfillment for the award of the degree of")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("BACHELOR OF TECHNOLOGY")
r.font.name = "Times New Roman"
r.font.size = Pt(15)
r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("in\nINFORMATION TECHNOLOGY")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.font.bold = True

# Logo
logo_path = os.path.join(ASSETS_DIR, "ride_p1_img1_808x241.png")
if os.path.exists(logo_path):
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(4)
    p_logo.paragraph_format.space_after = Pt(6)
    p_logo.add_run().add_picture(logo_path, width=Inches(3.2))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("DEPARTMENT OF INFORMATION TECHNOLOGY,\nPARUL INSTITUTE OF ENGINEERING AND TECHNOLOGY,\nPARUL UNIVERSITY,\nVADODARA, GUJARAT")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Ride Buddy")
r.font.name = "Times New Roman"
r.font.size = Pt(22)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(12)
r = p.add_run("Smart Intra-Campus Peer-to-Peer Ride Sharing Platform")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Submitted by:")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.bold = True

students = [
    "MR. AKSHAR PATEL (2303031080220)",
    "MR. OM DOBARIYA (2303031080039)",
    "MR. DIVYRAJSINH THAKOR (2303031080249)",
    "MR. DHRUVANSH MANIYA (2303031080212)"
]
for s in students:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(s)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Under the Guidance of")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Assistant Professor MS. HIMANI PARMAR")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("[2025-2026]")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

# ==========================================
# 2. CERTIFICATE PAGE
# ==========================================
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(16)
r = p.add_run("CERTIFICATE")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

cert_text = (
    'This is to certify that the Project Report entitled, "Ride Buddy" submitted by '
    '"Akshar Patel, Om Dobariya, Divyrajsinh Thakor and Dhruvansh Maniya" to Parul University, '
    'Vadodara, Gujarat, is a record of bonafide Project work carried out by them under my supervision '
    'and guidance, and is worthy of consideration for the award of the degree of Bachelor of Technology '
    'in Information Technology of the University.'
)
add_body_p(doc, cert_text, space_after=18)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Date:  ____________________\nPlace: Vadodara, Gujarat")
r.font.name = "Times New Roman"
r.font.size = Pt(11)

sig_table = doc.add_table(rows=2, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
for row in sig_table.rows:
    row.cells[0].width = Inches(3.5)
    row.cells[1].width = Inches(3.5)

cell_00 = sig_table.rows[0].cells[0]
p = cell_00.paragraphs[0]
p.paragraph_format.space_before = Pt(32)
r = p.add_run("______________________________\nSupervisor\nMs. Himani Parmar\nAssistant Professor")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

cell_01 = sig_table.rows[0].cells[1]
p = cell_01.paragraphs[0]
p.paragraph_format.space_before = Pt(32)
r = p.add_run("______________________________\nProject Coordinator\nMs. Dhenuka Patel\nAssistant Professor")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

cell_10 = sig_table.rows[1].cells[0]
p = cell_10.paragraphs[0]
p.paragraph_format.space_before = Pt(36)
r = p.add_run("______________________________\nHead, Dept. of Information Technology\nDr. Pooja Sapra")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

cell_11 = sig_table.rows[1].cells[1]
p = cell_11.paragraphs[0]
p.paragraph_format.space_before = Pt(36)
r = p.add_run("______________________________\nExternal Supervisor\nName:\nDesignation:")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

# ==========================================
# 3. ACKNOWLEDGEMENT PAGE
# ==========================================
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(16)
r = p.add_run("Acknowledgement")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

add_body_p(doc, "Behind any major work undertaken by an individual there lies the contribution of the people who helped him to cross all the hurdles to achieve his goal.", space_after=12)

add_body_p(doc, 
    "It gives me the immense pleasure to express my sense of sincere gratitude towards my respected guide "
    "(Assistant Professor) Prof. Himani Parmar for their persistent, outstanding, invaluable co-operation "
    "and guidance. It is my achievement to be guided under them. They are a constant source of encouragement "
    "and momentum that any intricacy becomes simple. I gained a lot of invaluable guidance and prompt suggestions "
    "from them during entire project work. I will be indebted to them forever and I take pride in working under them.",
    space_after=12)

add_body_p(doc,
    "We also express our deep sense of regards and thanks to Prof. Dr. Pooja Sapra, (Associate Professor) "
    "and Head of INFORMATION TECHNOLOGY Engineering Department. I feel very privileged to have had their "
    "precious advice, guidance, and leadership.",
    space_after=20)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Submitted By: -\n")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

for s in ["Akshar Patel (2303031080220)", "Om Dobariya (2303031080039)", "Divyrajsinh Thakor (2303031080249)", "Dhruvansh Maniya (2303031080212)"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(s)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True

# ==========================================
# 4. ABSTRACT PAGE
# ==========================================
doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(16)
r = p.add_run("ABSTRACT")
r.font.name = "Times New Roman"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

add_body_p(doc, 
    "RideBuddy is a web based ride-sharing application designed specifically for university students, enabling "
    "safe and affordable intra-city transportation from campus. Built to address the daily commuting challenges "
    "faced by students at Parul University, the platform connects student passengers with verified student riders "
    "through a seamless, real-time booking experience.", space_after=10)

add_body_p(doc,
    "The system operates on a predefined route model, offering popular destinations such as Waghodia Road, "
    "Vadodara Railway Station, Alkapuri, and Manjalpur allowing students to book rides with a single tap. "
    "Riders can register through a structured application process that includes vehicle and license verification, "
    "ensuring safety and accountability within the campus community.", space_after=10)

add_body_p(doc,
    "Key features include role-based access control (admin, rider, and user roles), real-time ride tracking with "
    "live map integration using Leaflet and OSRM routing, a rider dashboard for managing ride requests and earnings, "
    "and an admin panel for managing user verification and rider applications. Authentication is restricted to "
    "university email domains (@paruluniversity.ac.in), reinforcing the closed-campus ecosystem.", space_after=10)

add_body_p(doc,
    "The platform is built using modern web technologies including React, TypeScript, and Tailwind CSS on the frontend, "
    "with Supabase Cloud providing backend services for database management, authentication, file storage, and serverless "
    "edge functions. Real-time updates are powered through database subscriptions, ensuring both riders and passengers "
    "stay informed throughout the ride lifecycle.", space_after=10)

add_body_p(doc,
    "RideBuddy aims to promote carpooling culture within university campuses, reduce transportation costs for students, "
    "and provide a trusted, student-only ride sharing ecosystem.", space_after=14)

p_kw = doc.add_paragraph()
p_kw.paragraph_format.space_before = Pt(8)
r_k = p_kw.add_run("Keywords: ")
r_k.font.name = "Times New Roman"
r_k.font.size = Pt(11)
r_k.font.bold = True
r_v = p_kw.add_run("Ride Sharing Platform, Student Carpooling, University Email Authentication, Intra-City Student Mobility, User-Friendly")
r_v.font.name = "Times New Roman"
r_v.font.size = Pt(11)
r_v.font.italic = True

# Save progress
doc.save(OUTPUT_DOCX)
print("Saved prelims to", OUTPUT_DOCX)
